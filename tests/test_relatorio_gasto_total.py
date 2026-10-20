"""Relatório "Gasto total": precisa fechar exatamente com o card do painel."""
import io
from datetime import timedelta

import pytest
from openpyxl import load_workbook

from extensions import db
from services.tempo import hoje as data_de_hoje


@pytest.fixture()
def base(logado, app):
    """Cadastros mínimos criados direto no banco (não depende da API de peças)."""
    from models import Fornecedor, Peca
    v = logado.post("/api/veiculos", json={"prefixo": "FR-001", "placa": "ABC1D23",
                                           "tipo": "Ônibus"}).get_json()
    with app.app_context():
        oficina = Fornecedor(nome="Oficina Central", tipo="Oficina")
        peca = Peca(codigo="FIL-001", descricao="Filtro de óleo", grupo="Motor", quantidade=20,
                    custo_unitario=50)
        db.session.add_all([oficina, peca])
        db.session.commit()
        return {"veiculo": v, "oficina": {"id": oficina.id}, "peca": {"id": peca.id}}


@pytest.fixture()
def com_gastos(logado, base, app):
    """Um lançamento em cada categoria que compõe o Gasto total."""
    from models import (Ferramenta, GrupoConsumo, ItemNotaFiscalUniforme, ItemOS, ItemUniforme,
                        Lavagem, MovimentoEstoque, NotaFiscalUniforme, Orcamento, ServicoTerceiro,
                        TransferenciaFerramenta)
    hoje = data_de_hoje()
    vid = base["veiculo"]["id"]

    # Manutenção (OS): mão de obra 300 + serviços 100 + 2 peças de R$ 50
    ordem = logado.post("/api/ordens", json={"veiculo_id": vid, "tipo": "Corretiva", "grupo": "Freios",
                                             "custo_mao_obra": 300, "custo_servicos": 100}).get_json()

    with app.app_context():
        db.session.add(ItemOS(ordem_servico_id=ordem["id"], peca_id=base["peca"]["id"],
                              quantidade=2, valor_unitario=50, tipo_item="peça"))
        db.session.add(ServicoTerceiro(data=hoje, veiculo_id=vid, prestador="Retífica Sul",
                                       descricao="Retífica do cabeçote", valor=1500.50,
                                       categoria="Serviço de terceiros"))
        db.session.add(Lavagem(data=hoje, veiculo_id=vid, valor=80, observacao="Lavagem completa"))

        item = ItemUniforme(codigo="UNI-1", descricao="Camisa")
        db.session.add(item)
        db.session.flush()
        nota = NotaFiscalUniforme(numero="123", data_emissao=hoje, data_entrada=hoje,
                                  fornecedor_id=base["oficina"]["id"], status="Finalizada")
        nota.itens.append(ItemNotaFiscalUniforme(item_uniforme_id=item.id, tamanho="M",
                                                 quantidade=10, valor_unitario=35))
        db.session.add(nota)

        grupo = GrupoConsumo.query.filter_by(nome="Limpeza").first()
        if not grupo:
            grupo = GrupoConsumo(nome="Limpeza")
            db.session.add(grupo)
            db.session.flush()
        db.session.add(MovimentoEstoque(data=hoje, peca_id=base["peca"]["id"], tipo="saida",
                                        quantidade=3, custo_unitario=50, grupo_consumo_id=grupo.id))
        db.session.add(Orcamento(ano=hoje.year, mes=hoje.month, categoria="Geral", meta_valor=10000))

        ferr = Ferramenta(codigo="FER-1", descricao="Chave de roda")
        db.session.add(ferr)
        db.session.flush()
        saida = MovimentoEstoque(data=hoje, peca_id=base["peca"]["id"], tipo="saida",
                                 quantidade=2, custo_unitario=50, documento="Transf. ferramenta")
        db.session.add(saida)
        db.session.flush()
        db.session.add(TransferenciaFerramenta(data=hoje, peca_id=base["peca"]["id"],
                                               ferramenta_id=ferr.id, movimento_estoque_id=saida.id,
                                               quantidade=2, custo_unitario=50))
        db.session.commit()
    return base


def _periodo():
    hoje = data_de_hoje()
    return f"inicio={(hoje - timedelta(days=5)).isoformat()}&fim={hoje.isoformat()}"


def test_pdf_e_excel_fecham_com_o_card_do_painel(logado, com_gastos):
    q = _periodo()
    painel = logado.get(f"/api/painel/resumo?{q}").get_json()
    assert painel["gasto_total"] > 0
    # cada categoria tem valor: o teste só vale se todas estiverem preenchidas
    for campo in ("gasto_manutencao", "gasto_lavagem", "gasto_uniformes",
                  "gasto_grupos_consumo", "gasto_ferramentas_mecanicos"):
        assert painel[campo] > 0, campo

    pdf = logado.get(f"/relatorios/gasto-total.pdf?{q}")
    assert pdf.status_code == 200 and pdf.mimetype == "application/pdf"
    assert pdf.data.startswith(b"%PDF")

    xl = logado.get(f"/relatorios/gasto-total.xlsx?{q}")
    assert xl.status_code == 200
    wb = load_workbook(io.BytesIO(xl.data))
    assert wb.sheetnames[0] == "Resumo"
    ws = wb["Resumo"]
    valores = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=3).value
               for r in range(5, 12) if ws.cell(row=r, column=1).value}
    assert valores["Manutenção (OS)"] + valores["Serviços de terceiros"] == pytest.approx(
        painel["gasto_manutencao"])
    assert valores["Lavagem"] == pytest.approx(painel["gasto_lavagem"])
    assert valores["Uniformes (NF)"] == pytest.approx(painel["gasto_uniformes"])
    assert valores["Grupos de consumo"] == pytest.approx(painel["gasto_grupos_consumo"])
    assert valores["Ferramentas dos mecânicos"] == pytest.approx(painel["gasto_ferramentas_mecanicos"])
    # A soma das categorias (a fórmula do total) é o gasto total do painel.
    soma = sum(v for k, v in valores.items() if k != "GASTO TOTAL")
    assert soma == pytest.approx(painel["gasto_total"], abs=0.01)


def test_servico_dados_detalha_cada_categoria(app, com_gastos):
    from services import gasto_total
    hoje = data_de_hoje()
    with app.app_context():
        d = gasto_total.montar(hoje - timedelta(days=5), hoje)
    assert len(d["det_os"]) == 1 and d["det_os"][0][-1] == 500.0
    assert d["det_terceiros"][0][-1] == 1500.5
    assert d["det_lavagem"][0][-1] == 80.0
    assert d["det_uniformes"][0][-1] == 350.0
    assert d["det_ferramentas"][0][-2] == 100.0 or d["det_ferramentas"][0][-3] == 100.0
    assert d["gasto_total"] == pytest.approx(sum(x["valor"] for x in d["resumo"]), abs=0.01)
    assert d["aderencia"] is not None   # há meta cadastrada


def test_sem_meta_e_sem_lancamentos(logado, base):
    q = _periodo()
    assert logado.get(f"/relatorios/gasto-total.pdf?{q}").status_code == 200
    assert logado.get(f"/relatorios/gasto-total.xlsx?{q}").status_code == 200


def test_exige_login(cliente):
    assert cliente.get("/relatorios/gasto-total.pdf").status_code == 401
    assert cliente.get("/relatorios/gasto-total.xlsx").status_code == 401
