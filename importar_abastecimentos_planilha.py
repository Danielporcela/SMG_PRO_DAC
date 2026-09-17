# -*- coding: utf-8 -*-
"""Importa os abastecimentos da planilha "ABASTECIMENTO" para o banco do SGMF Pro DAC.

Lê a aba `LANÇAM` (um abastecimento por linha: VEÍCULO, DATA, LITROS, KM,
KM PERC, KM/L) e a aba `PLACA` (mapeia o número do veículo da planilha para
a placa) e grava um `Abastecimento` para cada linha válida, associado ao
veículo já cadastrado no sistema (localizado pela placa).

Nada é gravado pela metade: por padrão o script roda em modo "só
conferência" (dry-run) e mostra um resumo do que seria importado. Só grava
de verdade com a flag --commit. É seguro rodar de novo: linhas que já têm
um abastecimento igual (mesmo veículo + mesma data + mesmo KM) são puladas.

Uso (na raiz do projeto, mesma pasta do app.py):

    python importar_abastecimentos_planilha.py caminho/da/planilha.xlsx
    python importar_abastecimentos_planilha.py caminho/da/planilha.xlsx --commit

Requisitos: `pip install openpyxl` (se ainda não estiver instalado).
"""
import argparse
import sys
from pathlib import Path

from openpyxl import load_workbook

from extensions import db
from models import Abastecimento, Veiculo
from services.calculos import recalcular_abastecimento

ABA_LANCAMENTOS = "LANÇAM"
ABA_PLACAS = "PLACA"


def obter_app():
    """Funciona tanto com `app = Flask(...)` no app.py quanto com fábrica."""
    import app as modulo_app
    if hasattr(modulo_app, "app"):
        return modulo_app.app
    if hasattr(modulo_app, "create_app"):
        return modulo_app.create_app()
    if hasattr(modulo_app, "criar_app"):
        return modulo_app.criar_app()
    raise SystemExit("Não encontrei 'app' nem uma função de fábrica no app.py — "
                     "ajuste a função obter_app() deste script.")


def mapa_frota_para_placa(wb):
    """Lê a aba PLACA e devolve {número do veículo na planilha: placa}."""
    if ABA_PLACAS not in wb.sheetnames:
        print(f"  Aviso: aba '{ABA_PLACAS}' não encontrada na planilha — "
              "nenhum veículo poderá ser associado.")
        return {}
    ws = wb[ABA_PLACAS]
    mapa = {}
    for linha in ws.iter_rows(values_only=True):
        placa, frota = linha[0], linha[1]
        if placa and frota is not None:
            mapa[int(frota)] = str(placa).strip().upper()
    return mapa


def ler_lancamentos(wb):
    """Lê a aba LANÇAM e devolve uma lista de dicts com os dados brutos."""
    if ABA_LANCAMENTOS not in wb.sheetnames:
        raise SystemExit(f"Não encontrei a aba '{ABA_LANCAMENTOS}' na planilha.")
    ws = wb[ABA_LANCAMENTOS]
    linhas = list(ws.iter_rows(values_only=True))
    cabecalho = [str(c).strip().upper() if c else "" for c in linhas[0]]
    idx = {nome: cabecalho.index(nome) for nome in
           ("VEÍCULO", "DATA", "LITROS", "KM", "KM PERC", "KM/L") if nome in cabecalho}

    registros = []
    for linha in linhas[1:]:
        if linha[idx["VEÍCULO"]] is None:
            continue
        litros = linha[idx["LITROS"]] if "LITROS" in idx else None
        km = linha[idx["KM"]] if "KM" in idx else None
        data = linha[idx["DATA"]] if "DATA" in idx else None
        if litros is None or km is None or data is None:
            continue  # linha sem abastecimento lançado naquele dia
        registros.append({
            "veiculo_num": int(linha[idx["VEÍCULO"]]),
            "data": data.date() if hasattr(data, "date") else data,
            "litros": float(litros),
            "km_atual": float(km),
        })
    return registros


def localizar_planilha(caminho_planilha=None):
    """Localiza automaticamente a planilha de abastecimento.

    Se o caminho informado existir, ele é usado normalmente. Caso o arquivo
    tenha sido renomeado (por exemplo, acrescentando "(1)"), o script procura
    automaticamente na pasta do próprio script por um .xlsx que contenha
    ABASTECIMENTO, SETEMBRO e 2026 no nome e usa o mais recentemente alterado.
    """
    pasta_script = Path(__file__).resolve().parent

    if caminho_planilha:
        caminho = Path(caminho_planilha).expanduser()
        if caminho.exists() and caminho.is_file():
            return caminho

        # Se o caminho informado não existe, tenta localizar automaticamente.
        print(f"Aviso: arquivo informado não encontrado: {caminho}")

    candidatos = [
        arquivo for arquivo in pasta_script.glob("*.xlsx")
        if "ABASTECIMENTO" in arquivo.name.upper()
        and "SETEMBRO" in arquivo.name.upper()
        and "2026" in arquivo.name.upper()
    ]

    if not candidatos:
        # Fallback: qualquer planilha de abastecimento, caso o mês/ano
        # tenham sido alterados no nome do arquivo.
        candidatos = [
            arquivo for arquivo in pasta_script.glob("*.xlsx")
            if "ABASTECIMENTO" in arquivo.name.upper()
        ]

    if not candidatos:
        raise FileNotFoundError(
            "Nenhuma planilha .xlsx de abastecimento foi encontrada na pasta "
            f"do script: {pasta_script}"
        )

    escolhido = max(candidatos, key=lambda arquivo: arquivo.stat().st_mtime)
    print(f"Planilha localizada automaticamente: {escolhido.name}")
    return escolhido


def importar(caminho_planilha=None, commit=False):
    caminho_planilha = localizar_planilha(caminho_planilha)

    app = obter_app()
    with app.app_context():
        print(f"Lendo planilha: {caminho_planilha}")
        wb = load_workbook(caminho_planilha, data_only=True, read_only=True)

        frota_para_placa = mapa_frota_para_placa(wb)
        registros = ler_lancamentos(wb)
        registros.sort(key=lambda r: (r["veiculo_num"], r["data"]))
        print(f"  {len(registros)} lançamentos com abastecimento na planilha.")

        # cache de veículos do sistema por placa e por prefixo (fallback
        # quando o número da frota não tem placa mapeada na planilha)
        veiculos_por_placa = {v.placa.strip().upper(): v for v in Veiculo.query.all() if v.placa}
        veiculos_por_prefixo = {v.prefixo.strip().upper(): v for v in Veiculo.query.all() if v.prefixo}

        importados, duplicados, sem_veiculo, sem_mapa = [], [], [], set()

        for reg in registros:
            placa = frota_para_placa.get(reg["veiculo_num"])
            veiculo = veiculos_por_placa.get(placa) if placa else None
            if not veiculo:
                # fallback: tenta casar pelo prefixo do veículo == número da frota
                veiculo = veiculos_por_prefixo.get(str(reg["veiculo_num"]))
            if not veiculo:
                if placa:
                    sem_veiculo.append((reg["veiculo_num"], placa))
                else:
                    sem_mapa.add(reg["veiculo_num"])
                continue

            existente = Abastecimento.query.filter_by(
                veiculo_id=veiculo.id, data=reg["data"], km_atual=reg["km_atual"]
            ).first()
            if existente:
                duplicados.append(reg)
                continue

            abast = Abastecimento(
                data=reg["data"], veiculo_id=veiculo.id,
                combustivel=veiculo.combustivel, km_atual=reg["km_atual"],
                litros=reg["litros"], tanque_cheio=True,
            )
            db.session.add(abast)
            db.session.flush()  # garante abast.id para o cálculo do "anterior"
            recalcular_abastecimento(abast)
            importados.append(abast)

        print("\nResumo:")
        print(f"  A importar: {len(importados)}")
        print(f"  Já existentes (puladas): {len(duplicados)}")
        if sem_veiculo:
            faltantes = sorted(set(sem_veiculo))
            print(f"  Sem veículo cadastrado no sistema para a placa (veículo nº · placa): "
                  f"{faltantes}")
        if sem_mapa:
            print(f"  Sem placa mapeada na aba '{ABA_PLACAS}' para os veículos nº: "
                  f"{sorted(sem_mapa)}")

        if not commit:
            db.session.rollback()
            print("\nModo conferência (dry-run) — nada foi gravado. "
                  "Rode de novo com --commit para gravar de verdade.")
            return

        try:
            db.session.commit()
        except Exception:
            db.session.rollback()
            raise
        print(f"\n✓ {len(importados)} abastecimento(s) gravado(s) no banco.")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "planilha",
        nargs="?",
        help="Caminho do arquivo .xlsx (opcional; se omitido, será localizado automaticamente)"
    )
    parser.add_argument(
        "--commit",
        action="store_true",
        help="Grava de verdade (sem isso, roda só em modo conferência)"
    )
    args = parser.parse_args()
    importar(args.planilha, commit=args.commit)
    sys.exit(0)
